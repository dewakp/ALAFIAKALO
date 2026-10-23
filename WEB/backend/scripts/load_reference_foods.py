# Copyright © 2026 Wole Akpose / 6igma Health Inc.
# All rights reserved. ALAFIA — proprietary and confidential.

"""Load a food-composition authority into `reference_foods` (migration av001).

WAFCT 2019 — FAO/INFOODS West African Food Composition Table — is the first.
It holds what USDA does not: Vernonia (bitter leaf), melon seed (egusi),
dikanut (ogbono), toasted grated cassava (gari), papaya.

⚠️ LICENCE. WAFCT 2019 is CC BY-NC-SA 3.0 IGO (non-commercial), and ALAFIA is a
paid product. Operator, 2026-10-23: *"build with it. I am asking FAO for
permission."* So NOTHING from the table is in the repository or the image: the
operator supplies the workbook, and loading PRODUCTION is the step that waits
for FAO's answer. Download: https://www.fao.org/fileadmin/user_upload/faoweb/2020/WAFCT_2019.xlsx

Two steps, because the workbook reader is not a production dependency:

  extract   reads the .xlsx (needs openpyxl) -> one JSON file, outside the repo
            docker run --rm -v "$DIR":/w python:3.12-slim sh -c \\
              "pip install -q openpyxl && python -I /w/load_reference_foods.py extract \\
               --xlsx /w/WAFCT_2019.xlsx --out /w/wafct_2019.json"

  load      writes that JSON into the database; DRY RUN unless --apply
            python scripts/load_reference_foods.py load --json /w/wafct_2019.json [--apply]

Mapping decisions, each a place a naive copy is wrong:
  * Sheet "05 NV_sum_57" — per 100 g EDIBLE PORTION, 57 components; columns are
    found by their INFOODS TAGNAME (row 3), never by position.
  * carbs_g = CHOAVLDF + FIBTG. WAFCT prints AVAILABLE carbohydrate, which
    excludes fibre; ALAFIA's carbs_g follows USDA 1005 ("by difference"),
    which includes it. Copying CHOAVLDF alone reads every WAFCT food low.
  * vitamin_a_iu = RETOL/0.3 + CARTBEQ/0.6 — USDA's own IU convention
    (1 IU = 0.3 µg retinol = 0.6 µg β-carotene; CARTBEQ already folds in
    α-carotene and cryptoxanthin at half weight). There is no single RAE->IU
    factor, so VITA_RAE is not converted.
  * vitamin_d_iu = VITD µg × 40 (exact by definition).
  * "[1.2]" is a value WAFCT imputed or borrowed: kept, and its key listed in
    `imputed`. "tr" (trace) is 0.0 and also listed. A blank is ABSENT — never 0.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

SOURCE = "wafct_2019"
SHEET = "05 NV_sum_57 (per 100g EP)"

# INFOODS tagname -> (ALAFIA key, multiplier)
DIRECT = {
    "ENERC_KCAL": ("calories", 1.0), "PROTCNT": ("protein_g", 1.0), "FAT": ("fat_g", 1.0),
    "FIBTG": ("fiber_g", 1.0), "WATER": ("water_ml", 1.0), "ALC": ("alcohol_g", 1.0),
    "CA": ("calcium_mg", 1.0), "FE": ("iron_mg", 1.0), "MG": ("magnesium_mg", 1.0),
    "P": ("phosphorus_mg", 1.0), "K": ("potassium_mg", 1.0), "NA": ("sodium_mg", 1.0),
    "ZN": ("zinc_mg", 1.0), "CU": ("copper_mg", 1.0), "VITE": ("vitamin_e_mg", 1.0),
    "THIA": ("vitamin_b1_thiamine_mg", 1.0), "RIBF": ("vitamin_b2_riboflavin_mg", 1.0),
    "NIA": ("vitamin_b3_niacin_mg", 1.0), "VITB6C": ("vitamin_b6_mg", 1.0),
    "FOL": ("vitamin_b9_folate_mcg", 1.0), "VITB12": ("vitamin_b12_mcg", 1.0),
    "VITC": ("vitamin_c_mg", 1.0), "CHOLE": ("cholesterol_mg", 1.0),
    "FASAT": ("saturated_fat_g", 1.0), "FAMS": ("monounsaturated_fat_g", 1.0),
    "FAPU": ("polyunsaturated_fat_g", 1.0), "VITD": ("vitamin_d_iu", 40.0),
    "PHYTCPP": ("phytate_mg", 1.0),
}

# ── PhyFoodComp 1.0 (FAO/INFOODS/IZiNCG, phytate) ─────────────────────────────
# Phytate only — it ADDS a component to a food resolved elsewhere and is never
# offered as a match for the food's whole profile. "FAO/IZiNCG © February 2018",
# no Creative Commons grant: the same permission gate as WAFCT.
PHY_SOURCE = "phyfoodcomp_1_0"
# Most specific measurement first. HPLC/HPAE separates inositol hexaphosphate;
# precipitation methods also catch lower inositol phosphates and read high.
PHY_METHODS = [("IP6(mg)", "HPLC IP6"), ("PHYTCPP(mg)", "anion exchange"),
               ("PHYTCA (mg)", "K-PHYT kit"), ("PHYTCPPD(mg)", "direct precipitation"),
               ("PHYTCPPI(mg)", "indirect precipitation"), ("PHYTC-(mg)", "colorimetry"),
               ("PHYT-", "unknown method")]
PHY_PP = ["PPD(mg)", "PPI(mg)", "PP-(mg)"]
# Phosphorus mass fraction of phytic acid, C6H18O24P6: 6 x 30.974 / 660.04.
PHYTATE_P_FRACTION = 0.2816


def _cell(v):
    """-> (value or None, imputed?, trace?)"""
    if v is None:
        return None, False, False
    s = str(v).strip()
    if not s or s in ("-", "—", "None"):
        return None, False, False
    if s.lower() == "tr":
        return 0.0, False, True
    imputed = s.startswith("[") and s.endswith("]")
    s = s.strip("[]").replace(",", ".")
    try:
        return float(s), imputed, False
    except ValueError:
        return None, False, False


def _lod(v):
    """PhyFoodComp prints "< LOD" for below detection: present at ~0."""
    if v is not None and "lod" in str(v).lower():
        return 0.0, True
    val, imp, tr = _cell(v)
    return val, (imp or tr)


def extract_phyfood(xlsx: Path) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    foods, skipped = [], 0
    for ws in wb.worksheets:
        if not re.match(r"^\d{2} ", ws.title):
            continue
        rows = list(ws.iter_rows(values_only=True))
        col = {str(h).strip(): i for i, h in enumerate(rows[0]) if h}
        def g(r, name):
            # Exact header first, else the one header that STARTS with it —
            # several are long ("Processing / Influencing factors (…)").
            i = col.get(name)
            if i is None:
                hits = [j for h, j in col.items() if h.startswith(name)]
                i = hits[0] if len(hits) == 1 else None
            return r[i] if i is not None and i < len(r) else None
        for r in rows[2:]:
            code, name = g(r, "Food item ID"), g(r, "Food name in English")
            if not code or not name:
                skipped += 1
                continue
            phy, method, flagged = None, None, []
            for header, label in PHY_METHODS:
                v, f = _lod(g(r, header))
                if v is not None:
                    phy, method = v, label
                    if f:
                        flagged.append("phytate_mg")
                    break
            if phy is None:
                skipped += 1
                continue
            pp = None
            for header in PHY_PP:
                v, f = _lod(g(r, header))
                if v is not None:
                    pp = v
                    if f:
                        flagged.append("phytate_phosphorus_mg")
                    break
            if pp is None:
                pp = round(phy * PHYTATE_P_FRACTION, 4)
            nutrients = {"phytate_mg": round(phy, 4), "phytate_phosphorus_mg": round(pp, 4)}
            context = "; ".join(str(x) for x in (
                g(r, "Country, region"), g(r, "Processing / Influencing"),
                f"method: {method}", g(r, "Biblioid")) if x)
            foods.append({
                "source": PHY_SOURCE, "code": str(code).strip(), "name_en": str(name).strip()[:400],
                "name_fr": None, "scientific_name": (str(g(r, "Species/Subspecies")).strip()[:200]
                                                     if g(r, "Species/Subspecies") else None),
                "food_group": ws.title[3:].strip()[:200], "nutrients": nutrients,
                "imputed": flagged or None, "biblio": context or None,
            })
    print(f"PhyFoodComp: {len(foods)} entries with a phytate value ({skipped} rows without one)")
    return foods


def extract(xlsx: Path, out: Path) -> None:
    if "phyfood" in xlsx.name.lower():
        foods = extract_phyfood(xlsx)
        out.write_text(json.dumps(foods, ensure_ascii=False))
        print(f"-> {out}")
        return
    import openpyxl

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    rows = list(wb[SHEET].iter_rows(values_only=True))
    header, tags = rows[0], rows[2]
    col = {}
    for i, (h, t) in enumerate(zip(header, tags)):
        tag = (str(t).strip() if t else "")
        if tag == "ENERC":  # two ENERC columns: kJ then kcal — told apart by the header
            tag = "ENERC_KCAL" if "kcal" in str(h).lower() else "ENERC_KJ"
        if tag:
            col.setdefault(tag, i)
    need = {"ENERC_KCAL", "PROTCNT", "FAT", "CHOAVLDF", "FIBTG", "RETOL", "CARTBEQ"}
    missing = need - set(col)
    if missing:
        sys.exit(f"workbook layout changed — tags not found: {sorted(missing)}")

    foods, group, rejected = [], None, 0
    for r in rows[3:]:
        code, name = r[0], r[1]
        if code and not name:  # a food-group heading row
            group = str(code).strip()
            continue
        if not code or not name or not re.match(r"^\d{2}_\d+", str(code)):
            rejected += 1
            continue
        nutrients, imputed = {}, []
        for tag, (key, mult) in DIRECT.items():
            if tag not in col:
                continue
            v, imp, tr = _cell(r[col[tag]])
            if v is None:
                continue
            nutrients[key] = round(v * mult, 4)
            if imp or tr:
                imputed.append(key)
        carb, ci, _ = _cell(r[col["CHOAVLDF"]])
        fib, fi, _ = _cell(r[col["FIBTG"]])
        if carb is not None:
            nutrients["carbs_g"] = round(carb + (fib or 0.0), 4)
            if ci or fi:
                imputed.append("carbs_g")
        ret, ri, _ = _cell(r[col["RETOL"]])
        car, cri, _ = _cell(r[col["CARTBEQ"]])
        if ret is not None or car is not None:
            nutrients["vitamin_a_iu"] = round((ret or 0.0) / 0.3 + (car or 0.0) / 0.6, 2)
            if ri or cri:
                imputed.append("vitamin_a_iu")
        if "calories" not in nutrients:
            rejected += 1
            continue
        foods.append({
            "source": SOURCE, "code": str(code).strip(), "name_en": str(name).strip(),
            "name_fr": (str(r[2]).strip() if r[2] else None),
            "scientific_name": (str(r[3]).strip() if r[3] else None),
            "food_group": group, "nutrients": nutrients,
            "imputed": sorted(set(imputed)) or None,
            "biblio": (str(r[4]).strip() if r[4] else None),
        })
    out.write_text(json.dumps(foods, ensure_ascii=False))
    print(f"extracted {len(foods)} foods ({rejected} rows skipped: headings, notes, no energy) -> {out}")


async def load(path: Path, apply: bool) -> None:
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from sqlalchemy import select

    from app.core.database import async_session
    from app.models.reference_food import ReferenceFood

    foods = json.loads(path.read_text())
    async with async_session() as db:
        existing = {(r.source, r.code): r for r in (await db.execute(select(ReferenceFood))).scalars()}
        new = [f for f in foods if (f["source"], f["code"]) not in existing]
        changed = [f for f in foods if (f["source"], f["code"]) in existing
                   and existing[(f["source"], f["code"])].nutrients != f["nutrients"]]
        print(f"{len(foods)} foods in file: {len(new)} new, {len(changed)} changed, "
              f"{len(foods) - len(new) - len(changed)} unchanged")
        if not apply:
            print("DRY RUN — nothing written. Re-run with --apply.")
            return
        for f in new:
            db.add(ReferenceFood(**f))
        for f in changed:
            row = existing[(f["source"], f["code"])]
            for k, v in f.items():
                setattr(row, k, v)
        await db.commit()
        print(f"wrote {len(new)} new, updated {len(changed)}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("extract")
    e.add_argument("--xlsx", required=True, type=Path)
    e.add_argument("--out", required=True, type=Path)
    lo = sub.add_parser("load")
    lo.add_argument("--json", required=True, type=Path)
    lo.add_argument("--apply", action="store_true")
    a = ap.parse_args()
    if a.cmd == "extract":
        extract(a.xlsx, a.out)
    else:
        import asyncio
        asyncio.run(load(a.json, a.apply))


if __name__ == "__main__":
    main()
